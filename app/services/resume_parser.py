"""Resume text extraction and structured field parsing.

The parser is deliberately conservative. It reports only what it actually found in the
document, attaches a confidence to each field, and records a warning when a section is
missing. It never substitutes a plausible value for an absent one, because the extracted
data goes on to populate application answers.

Confidence semantics:

* ``0.9`` — read from an explicitly labelled section ("SKILLS", "EDUCATION").
* ``0.6`` — matched by a specific pattern anywhere in the document (email, degree line).
* ``0.4`` — inferred by scanning the whole document against a known vocabulary.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.db.models import utcnow
from app.schemas.resume import ParsedResumeFields
from app.services.text import canonical_skill, canonical_skills, collapse_whitespace

_EMAIL = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
# Matches a run of digits with optional separators; the digit count is checked separately,
# because grouping varies widely ("+91 98765 43210", "(555) 123-4567", "9876543210").
_PHONE = re.compile(
    r"(?<![\w@.])(?:\+\d{1,3}[\s\-.]?)?(?:\(?\d{2,5}\)?[\s\-.]?){2,4}\d{2,5}(?![\w@])"
)
_URL = re.compile(
    r"(?:https?://|www\.)[^\s,;)\]]+"
    r"|(?:github|gitlab|linkedin|bitbucket|kaggle|medium|dev|hashnode|leetcode|stackoverflow|huggingface)"
    r"\.(?:com|io|to|org)/[^\s,;)\]|]+",
    re.IGNORECASE,
)
_YEAR = re.compile(r"(19|20)\d{2}")
_BULLET = re.compile(r"^\s*(?:[-*\u2022\u25cf\u25aa\u2023\u2043\u00b7>]|\d+[.)])\s+")

_MONTHS = (
    "jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec|"
    "january|february|march|april|june|july|august|september|october|november|december"
)
_DATE_RANGE = re.compile(
    rf"(?P<start_month>{_MONTHS})?\.?\s*(?P<start_year>(?:19|20)\d{{2}})\s*"
    r"(?:-|–|—|to|until|through)\s*"
    rf"(?:(?P<end_month>{_MONTHS})?\.?\s*(?P<end_year>(?:19|20)\d{{2}})|(?P<present>present|current|now|till date|ongoing))",
    re.IGNORECASE,
)

_DEGREE_PATTERN = re.compile(
    r"\b("
    r"b\.?\s?tech|bachelor(?:'s)?(?:\s+of\s+\w+)?|b\.?\s?e\b|b\.?\s?sc|b\.?\s?c\.?\s?a|bca|"
    r"m\.?\s?tech|master(?:'s)?(?:\s+of\s+\w+)?|m\.?\s?e\b|m\.?\s?sc|m\.?\s?c\.?\s?a|mca|mba|"
    r"ph\.?\s?d|doctorate|diploma|b\.?\s?com|m\.?\s?com|12th|10th|higher secondary|"
    r"intermediate|senior secondary"
    r")\b",
    re.IGNORECASE,
)

SECTION_PATTERNS: dict[str, tuple[str, ...]] = {
    "summary": ("summary", "objective", "profile", "about me", "career objective"),
    "skills": (
        "skills",
        "technical skills",
        "technical proficiency",
        "core competencies",
        "technologies",
        "tech stack",
        "areas of expertise",
        "skills & tools",
        "tools and technologies",
    ),
    "experience": (
        "experience",
        "work experience",
        "professional experience",
        "employment",
        "employment history",
        "work history",
        "career history",
    ),
    "internships": ("internship", "internships", "internship experience", "trainings"),
    "education": (
        "education",
        "academic",
        "academics",
        "academic qualifications",
        "qualifications",
    ),
    "projects": (
        "projects",
        "academic projects",
        "personal projects",
        "key projects",
        "project work",
    ),
    "certifications": (
        "certification",
        "certifications",
        "courses",
        "certificates",
        "licenses & certifications",
        "online courses",
    ),
    "achievements": (
        "achievement",
        "achievements",
        "awards",
        "honors",
        "honours",
        "accomplishments",
        "extra-curricular",
        "activities",
        "publications",
    ),
    "languages": ("languages", "language proficiency"),
    "interests": ("interests", "hobbies"),
}

KNOWN_SKILL_VOCABULARY: tuple[str, ...] = (
    "python",
    "java",
    "javascript",
    "typescript",
    "c",
    "c++",
    "c#",
    "go",
    "rust",
    "ruby",
    "php",
    "kotlin",
    "swift",
    "scala",
    "r",
    "matlab",
    "bash",
    "shell",
    "perl",
    "dart",
    "html",
    "css",
    "sass",
    "tailwind",
    "bootstrap",
    "react",
    "next.js",
    "vue",
    "angular",
    "svelte",
    "jquery",
    "redux",
    "react native",
    "flutter",
    "streamlit",
    "node.js",
    "express",
    "django",
    "flask",
    "fastapi",
    "spring",
    "spring boot",
    ".net",
    "laravel",
    "rails",
    "graphql",
    "rest api",
    "grpc",
    "websockets",
    "microservices",
    "sql",
    "postgresql",
    "mysql",
    "sqlite",
    "mongodb",
    "redis",
    "cassandra",
    "dynamodb",
    "elasticsearch",
    "neo4j",
    "firebase",
    "supabase",
    "sqlalchemy",
    "prisma",
    "aws",
    "azure",
    "google cloud",
    "docker",
    "kubernetes",
    "terraform",
    "ansible",
    "jenkins",
    "github actions",
    "gitlab ci",
    "cicd",
    "linux",
    "nginx",
    "kafka",
    "rabbitmq",
    "airflow",
    "prometheus",
    "grafana",
    "serverless",
    "lambda",
    "machine learning",
    "deep learning",
    "natural language processing",
    "computer vision",
    "generative ai",
    "llm",
    "retrieval augmented generation",
    "prompt engineering",
    "langchain",
    "llamaindex",
    "transformers",
    "hugging face",
    "pytorch",
    "tensorflow",
    "keras",
    "scikit-learn",
    "numpy",
    "pandas",
    "matplotlib",
    "seaborn",
    "opencv",
    "vector database",
    "pinecone",
    "chromadb",
    "faiss",
    "embeddings",
    "fine-tuning",
    "mlops",
    "mlflow",
    "agentic ai",
    "openai",
    "anthropic",
    "git",
    "github",
    "gitlab",
    "jira",
    "postman",
    "pytest",
    "junit",
    "selenium",
    "playwright",
    "cypress",
    "testing",
    "agile",
    "scrum",
    "data structures and algorithms",
    "object oriented programming",
    "operating systems",
    "computer networks",
    "dbms",
    "system design",
    "distributed systems",
    "excel",
    "power bi",
    "tableau",
    "figma",
)

_SOFT_SKILL_VOCABULARY: tuple[str, ...] = (
    "communication",
    "teamwork",
    "leadership",
    "problem solving",
    "critical thinking",
    "collaboration",
    "adaptability",
    "time management",
    "mentoring",
    "presentation",
    "stakeholder management",
    "ownership",
    "attention to detail",
)


class ResumeExtractionError(RuntimeError):
    """Raised when a file cannot be read as a resume at all."""


@dataclass(slots=True)
class ParseResult:
    fields: ParsedResumeFields
    warnings: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------------------
# Text extraction
# --------------------------------------------------------------------------------------


def extract_text_from_pdf(data: bytes) -> str:
    try:
        import fitz
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ResumeExtractionError("PyMuPDF is not installed; cannot read PDF files") from exc

    try:
        with fitz.open(stream=data, filetype="pdf") as document:
            if document.needs_pass:
                raise ResumeExtractionError(
                    "this PDF is password protected; remove the password and upload again"
                )
            pages = [page.get_text("text") for page in document]
    except ResumeExtractionError:
        raise
    except Exception as exc:
        raise ResumeExtractionError(f"could not read the PDF: {exc}") from exc
    return "\n".join(pages)


def extract_pdf_links(data: bytes) -> list[str]:
    """Hyperlink targets from a PDF, which are annotations rather than page text."""
    try:
        import fitz
    except ImportError:  # pragma: no cover - dependency is declared
        return []
    try:
        with fitz.open(stream=data, filetype="pdf") as document:
            if document.needs_pass:
                return []
            return _pdf_link_uris(document)
    except Exception:  # pragma: no cover - a readable PDF already parsed above
        return []


def _pdf_link_uris(document: Any) -> list[str]:
    """Collect hyperlink targets, which are annotations rather than page text.

    A resume whose header reads "Portfolio | LinkedIn | GitHub" carries the actual URLs in
    link annotations; the visible text is only the label. Reading text alone loses every
    link, which for a fresher is often the most valuable evidence on the page.

    These are returned separately rather than appended to the text on purpose: the skills
    section is frequently the last one on the page, so appended URLs were absorbed into it
    and "github.com", "in" and "u" became skills.
    """
    found: list[str] = []
    for page in document:
        try:
            links = page.get_links()
        except Exception:  # pragma: no cover - defensive, varies by PDF
            continue
        for link in links:
            uri = str(link.get("uri") or "").strip()
            if uri and uri.lower().startswith(("http://", "https://")) and uri not in found:
                found.append(uri)
    return found[:20]


def extract_text_from_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ResumeExtractionError("python-docx is not installed; cannot read DOCX files") from exc

    import io

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeExtractionError(f"could not read the DOCX file: {exc}") from exc

    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n".join(blocks)


#: A word split across lines by justified typesetting ("execu-\ntion"). Only a lowercase
#: continuation is joined, so a genuine compound broken at the hyphen ("Full-\nStack") keeps
#: its hyphen. LaTeX resumes hyphenate constantly, and an unjoined fragment becomes a
#: garbage skill ("terminal-based work" + "flows") that can never match a job requirement.
_LINE_BREAK_HYPHEN = re.compile(r"(\w)[-\u2010\u2011]\n([a-z])")


def dehyphenate(text: str) -> str:
    """Rejoin words broken across lines by hyphenation."""
    return _LINE_BREAK_HYPHEN.sub(r"\1\2", text)


def extract_text(data: bytes, extension: str) -> str:
    suffix = extension.lower()
    if suffix == ".pdf":
        return dehyphenate(extract_text_from_pdf(data))
    if suffix == ".docx":
        return dehyphenate(extract_text_from_docx(data))
    raise ResumeExtractionError(f"unsupported resume format: {suffix}")


def extract_text_from_path(path: str | Path) -> str:
    candidate = Path(path)
    return extract_text(candidate.read_bytes(), candidate.suffix)


# --------------------------------------------------------------------------------------
# Sectioning
# --------------------------------------------------------------------------------------


def _heading_key(line: str) -> str | None:
    """Return the section key if ``line`` looks like a section heading."""
    stripped = line.strip().strip(":").strip()
    if not stripped or len(stripped) > 60 or _BULLET.match(line):
        return None
    lowered = re.sub(r"[^a-z& ]+", " ", stripped.lower())
    lowered = collapse_whitespace(lowered)
    if not lowered:
        return None
    for key, aliases in SECTION_PATTERNS.items():
        for alias in aliases:
            if lowered == alias or lowered == alias + "s":
                return key
    # A short ALL-CAPS line that contains an alias is still a heading
    # ("TECHNICAL SKILLS & TOOLS"), but a sentence that merely mentions one is not.
    if stripped.isupper() and len(stripped.split()) <= 5:
        for key, aliases in SECTION_PATTERNS.items():
            if any(alias in lowered for alias in aliases):
                return key
    return None


def split_sections(text: str) -> dict[str, list[str]]:
    """Split resume text into labelled sections, keyed by :data:`SECTION_PATTERNS`."""
    sections: dict[str, list[str]] = {"_preamble": []}
    current = "_preamble"
    for raw_line in text.splitlines():
        key = _heading_key(raw_line)
        if key is not None:
            current = key
            sections.setdefault(current, [])
            continue
        if raw_line.strip():
            sections.setdefault(current, []).append(raw_line.rstrip())
    return sections


def _bullets(lines: list[str]) -> list[str]:
    """Group section lines into logical items.

    Two layouts occur in practice and they need opposite treatment. When a section uses
    bullet markers, every marker starts a new item and unmarked lines are wrapped
    continuations. When it does not, each line is its own entry, and only a line that
    clearly continues the previous one (starts lowercase or with a joining punctuation
    mark) is merged. Merging unconditionally is what collapses two education rows into one.
    """
    has_markers = any(_BULLET.match(line) for line in lines)
    items: list[str] = []
    buffer = ""

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if has_markers:
            if _BULLET.match(line):
                if buffer:
                    items.append(collapse_whitespace(buffer))
                buffer = _BULLET.sub("", line, count=1).strip()
            else:
                buffer = f"{buffer} {stripped}".strip()
            continue

        is_continuation = bool(buffer) and (
            stripped[0].islower()
            or stripped[0] in ",;:&/-"
            or buffer.rstrip().endswith((",", "-", "&"))
        )
        if is_continuation:
            buffer = f"{buffer} {stripped}"
        else:
            if buffer:
                items.append(collapse_whitespace(buffer))
            buffer = stripped

    if buffer:
        items.append(collapse_whitespace(buffer))
    return [item for item in items if item]


# --------------------------------------------------------------------------------------
# Field extraction
# --------------------------------------------------------------------------------------


def _split_skill_line(line: str) -> list[str]:
    body = re.sub(r"^[^:]{0,40}:", "", line) if ":" in line[:42] else line
    parts = re.split(r"[,;|/\u2022\u00b7]|\s{3,}", body)
    return [part.strip(" .-\t") for part in parts if part.strip(" .-\t")]


def _extract_skills(sections: dict[str, list[str]], full_text: str) -> tuple[list[str], float]:
    explicit: list[str] = []
    for line in sections.get("skills", []):
        explicit.extend(_split_skill_line(line))

    declared = canonical_skills(explicit)
    # Keep only tokens that look like skills, not sentences accidentally in the section.
    declared = [skill for skill in declared if len(skill.split()) <= 5]

    haystack = f" {re.sub(r'[^a-z0-9+#. ]+', ' ', full_text.lower())} "
    inferred = [
        skill
        for skill in KNOWN_SKILL_VOCABULARY
        if skill not in declared and f" {skill} " in haystack
    ]

    combined = canonical_skills(declared + inferred)
    confidence = 0.9 if declared else (0.4 if combined else 0.0)
    return combined, confidence


def _extract_soft_skills(full_text: str) -> list[str]:
    lowered = full_text.lower()
    return [skill for skill in _SOFT_SKILL_VOCABULARY if skill in lowered]


_INSTITUTION_PATTERN = re.compile(
    r"\b[A-Z][\w.&'\- ]{2,70}"
    r"(?:University|College|Institute|Institute of Technology|School|Academy|Polytechnic"
    r"|IIT|NIT|IIIT|BITS)\b"
)


def _education_block(lines: list[str], index: int) -> str:
    """The degree line plus its neighbours.

    A one-column resume prints an education entry as several separate lines — institution,
    city, degree, then the date range — so the year and the institution are almost never on
    the same line as the degree. Searching only the degree line loses the graduation year,
    which then blocks the graduation-date question on every application for no reason.
    """
    start = max(0, index - 3)
    end = min(len(lines), index + 4)
    return "\n".join(lines[start:end])


def _extract_education(sections: dict[str, list[str]], full_text: str) -> list[dict[str, str]]:
    lines = sections.get("education") or []
    scope = lines if lines else full_text.splitlines()
    entries: list[dict[str, str]] = []
    items = _bullets(list(scope))

    for index, item in enumerate(items):
        degree_match = _DEGREE_PATTERN.search(item)
        if degree_match is None:
            continue

        block = _education_block(items, index)
        year_values = re.findall(r"(?:19|20)\d{2}", item) or re.findall(r"(?:19|20)\d{2}", block)

        institution = ""
        institution_match = re.search(
            r"(?:at|from|,)\s+([A-Z][\w.&'\- ]{3,70}"
            r"(?:University|College|Institute|School|Academy|IIT|NIT|IIIT))",
            item,
        )
        if institution_match:
            institution = institution_match.group(1).strip()
        else:
            block_match = _INSTITUTION_PATTERN.search(block)
            if block_match:
                institution = collapse_whitespace(block_match.group(0))[:120]
            elif "," in item:
                tail = item.split(",", 1)[1]
                institution = collapse_whitespace(
                    re.sub(r"(?:19|20)\d{2}", "", tail).strip(" ,.-")
                )[:120]

        grade_match = re.search(
            r"(?:cgpa|gpa|percentage|score)\D{0,5}([\d.]{1,5})\s*%?", block, re.I
        )
        entries.append(
            {
                "degree": collapse_whitespace(degree_match.group(0)),
                "institution": institution,
                "field_of_study": "",
                "start_year": year_values[0] if len(year_values) > 1 else "",
                "end_year": year_values[-1] if year_values else "",
                "grade": grade_match.group(1) if grade_match else "",
                "raw": item[:300],
            }
        )
        if len(entries) >= 8:
            break
    return entries


def _looks_like_tech_list(line: str) -> bool:
    """True for a line that is a comma-separated technology list, not prose.

    A project in a one-column resume is a bare title line followed by its stack on the next
    line. Telling the two apart is what stops the stack being glued onto the project name.
    """
    stripped = line.strip().rstrip(".")
    if not stripped or len(stripped) > 140 or _BULLET.match(line):
        return False
    parts = [part.strip() for part in stripped.split(",") if part.strip()]
    if len(parts) < 2:
        return False
    if any(len(part.split()) > 4 for part in parts):
        return False
    known = sum(1 for part in parts if canonical_skill(part) in KNOWN_SKILL_VOCABULARY)
    return known >= max(1, len(parts) // 2)


def _looks_like_project_title(line: str) -> bool:
    """True for a project heading, false for a wrapped continuation of a bullet.

    Justified text wraps a bullet onto following unmarked lines, so "is it a bullet?" does
    not separate titles from prose. Two signals do: a resume project title is short and does
    not end in a sentence terminator, whereas a wrapped fragment almost always completes a
    sentence ("brand analysis, trend research and content creation.").
    """
    stripped = line.strip()
    if not stripped or len(stripped) > 70 or _BULLET.match(line):
        return False
    if stripped.endswith((".", ",", ";", ":")):
        return False
    return bool(re.match(r"[A-Z0-9]", stripped))


def _extract_projects(sections: dict[str, list[str]]) -> list[dict[str, object]]:
    """Group a projects section into one entry per project.

    Bullets and their wrapped continuations describe the project named by the most recent
    title line. Treating every line as its own project produced entries called "Built an
    autonomous LinkedIn AI agent using..." which then appeared verbatim as a project name in
    generated cover letters.
    """
    lines = sections.get("projects", [])
    projects: list[dict[str, object]] = []
    current: dict[str, object] | None = None

    def _flush() -> None:
        if current is not None:
            projects.append(current)

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue

        if _looks_like_tech_list(line) and current is not None and not current["technologies_raw"]:
            current["technologies_raw"] = collapse_whitespace(line)[:200]
            continue

        if _looks_like_project_title(line) and not _looks_like_tech_list(line):
            _flush()
            current = {
                "name": collapse_whitespace(line)[:120],
                "description": "",
                "technologies_raw": "",
                "url": "",
            }
            continue

        # The other common layout puts a whole project on one bulleted line as
        # "Name: description". Splitting it keeps that resume working, and a wrapped
        # continuation never reaches here because it has no title-shaped head.
        body = _BULLET.sub("", raw_line).strip()
        head, separator, tail = body.partition(" - ") if " - " in body else body.partition(":")
        if separator and tail.strip() and _looks_like_project_title(head):
            _flush()
            current = {
                "name": collapse_whitespace(head)[:120],
                "description": collapse_whitespace(tail)[:1200],
                "technologies_raw": "",
                "url": "",
            }
            continue

        if current is None:
            current = {
                "name": collapse_whitespace(body)[:120],
                "description": "",
                "technologies_raw": "",
                "url": "",
            }
            continue

        existing = str(current["description"])
        current["description"] = collapse_whitespace(f"{existing} {body}")[:1200]

    _flush()

    for project in projects:
        blob = (
            f"{project['name']} {project.pop('technologies_raw')} {project['description']}"
        ).lower()
        # Punctuation becomes a space so a comma-separated stack ("Python, FastAPI") still
        # matches; +, # and . survive because they are part of c++, c# and node.js.
        padded = " " + re.sub(r"[^a-z0-9+#.]+", " ", blob).strip() + " "
        project["technologies"] = [
            skill for skill in KNOWN_SKILL_VOCABULARY if f" {skill} " in padded
        ]
        url_match = _URL.search(str(project["description"]))
        project["url"] = url_match.group(0) if url_match else ""

    return [project for project in projects if project["name"]][:15]


def _extract_roles(lines: list[str]) -> list[dict[str, object]]:
    roles: list[dict[str, object]] = []
    for item in _bullets(lines):
        if len(item) < 6:
            continue
        date_match = _DATE_RANGE.search(item)
        roles.append(
            {
                "raw": item[:400],
                "start": (date_match.group("start_year") if date_match else ""),
                "end": (
                    "present"
                    if date_match and date_match.group("present")
                    else (date_match.group("end_year") if date_match else "")
                ),
            }
        )
        if len(roles) >= 15:
            break
    return roles


_MONTH_NUMBERS = {
    name: index
    for index, name in enumerate(
        ("jan,feb,mar,apr,may,jun,jul,aug,sep,oct,nov,dec".split(",")),
        start=1,
    )
}


def _month_number(name: str | None) -> int | None:
    if not name:
        return None
    return _MONTH_NUMBERS.get(name.strip().lower()[:3])


def _estimate_years_experience(sections: dict[str, list[str]], full_text: str) -> float | None:
    """Estimate professional experience from date ranges, or ``None`` if unknowable.

    An explicit "N years of experience" statement wins. Otherwise date ranges inside the
    experience section are summed. Internship and project sections are excluded because
    counting them would overstate a fresher's experience.
    """
    explicit = re.search(
        r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)\s+(?:of\s+)?(?:professional\s+|work\s+|industry\s+)?experience",
        full_text,
        re.IGNORECASE,
    )
    if explicit:
        try:
            return min(float(explicit.group(1)), 60.0)
        except ValueError:
            pass

    experience_lines = sections.get("experience") or []
    if not experience_lines:
        return None

    months = 0
    now = utcnow()
    for match in _DATE_RANGE.finditer("\n".join(experience_lines)):
        start_year = int(match.group("start_year"))
        start_month = _month_number(match.group("start_month")) or 1
        if match.group("present"):
            end_year, end_month = now.year, now.month
        else:
            end_year = int(match.group("end_year") or 0)
            end_month = _month_number(match.group("end_month")) or 12
        if end_year < start_year:
            continue
        # Month granularity matters most exactly where the number is smallest: a fresher
        # six months into a first job is 0.5 years, and rounding that to a whole year
        # either way changes which roles they are judged eligible for.
        span = (end_year - start_year) * 12 + (end_month - start_month)
        months += max(span, 1)
    if months == 0:
        return None
    return round(min(months / 12.0, 60.0), 1)


def _extract_graduation_date(education: list[dict[str, str]]) -> str | None:
    years = [entry.get("end_year") or "" for entry in education]
    numeric = sorted({int(year) for year in years if year.isdigit()})
    return str(numeric[-1]) if numeric else None


def _extract_contact(text: str) -> tuple[str, str, list[str]]:
    email_match = _EMAIL.search(text)
    email = email_match.group(0) if email_match else ""

    # Strip emails and URLs first: both contain digit runs that look like phone numbers.
    phone_scope = _URL.sub(" ", _EMAIL.sub(" ", text))
    phone = ""
    for match in _PHONE.finditer(phone_scope):
        digits = re.sub(r"\D", "", match.group(0))
        if 10 <= len(digits) <= 13:
            phone = collapse_whitespace(match.group(0))
            break

    links: list[str] = []
    for match in _URL.finditer(text):
        url = match.group(0).rstrip(".,);|")
        if not url.lower().startswith("http"):
            url = f"https://{url}"
        if url not in links:
            links.append(url)
    return email, phone, links[:10]


def _extract_name(text: str, email: str) -> str:
    """Take the first plausible name line from the header block.

    Resumes put the name on the first non-empty line far more often than anywhere else;
    anything containing an address, a URL or a digit is rejected rather than guessed at.
    """
    for line in text.splitlines()[:8]:
        candidate = collapse_whitespace(line)
        if not candidate or len(candidate) > 60:
            continue
        if any(char.isdigit() for char in candidate):
            continue
        if "@" in candidate or "http" in candidate.lower() or "|" in candidate:
            continue
        words = candidate.replace(".", "").split()
        if 1 < len(words) <= 5 and all(word[:1].isalpha() for word in words):
            if _heading_key(line) is None:
                return candidate
    if email:
        local = email.split("@", 1)[0]
        cleaned = re.sub(r"[._\-0-9]+", " ", local).strip()
        if len(cleaned.split()) >= 2:
            return cleaned.title()
    return ""


def parse_resume_text(text: str) -> ParseResult:
    """Parse raw resume text into structured fields with confidences and warnings."""
    warnings: list[str] = []
    if not text or not text.strip():
        return ParseResult(
            fields=ParsedResumeFields(),
            warnings=[
                "No readable text was found. The file may be a scanned image; this build "
                "has no OCR, so type the details into the profile screen instead."
            ],
        )
    if len(text.strip()) < 120:
        warnings.append(
            f"Only {len(text.strip())} characters of text were extracted, which is unusually "
            "short for a resume. Review every extracted field before accepting it."
        )

    sections = split_sections(text)
    email, phone, links = _extract_contact(text)
    skills, skills_confidence = _extract_skills(sections, text)
    education = _extract_education(sections, text)
    projects = _extract_projects(sections)
    experience = _extract_roles(sections.get("experience", []))
    internships = _extract_roles(sections.get("internships", []))
    certifications = [item[:200] for item in _bullets(sections.get("certifications", []))][:20]
    achievements = [item[:300] for item in _bullets(sections.get("achievements", []))][:20]
    years = _estimate_years_experience(sections, text)

    for key, label in (
        ("skills", "skills"),
        ("education", "education"),
        ("projects", "projects"),
    ):
        if not sections.get(key):
            warnings.append(
                f"No '{label}' section heading was recognised. Values shown for {label} were "
                "inferred from the whole document and need checking."
            )

    if not email:
        warnings.append("No email address was found in the resume.")
    if education and not any(entry.get("end_year") for entry in education):
        warnings.append("Education entries were found but none carried a year.")

    confidence = {
        "skills": skills_confidence,
        "education": 0.9 if sections.get("education") else (0.6 if education else 0.0),
        "projects": 0.9 if sections.get("projects") else 0.0,
        "experience": 0.9 if sections.get("experience") else 0.0,
        "certifications": 0.9 if sections.get("certifications") else 0.0,
        "achievements": 0.9 if sections.get("achievements") else 0.0,
        "email": 0.6 if email else 0.0,
        "phone": 0.6 if phone else 0.0,
        "full_name": 0.4,
        "years_experience": 0.6 if years is not None else 0.0,
    }

    fields = ParsedResumeFields(
        full_name=_extract_name(text, email),
        email=email,
        phone=phone,
        links=links,
        skills=skills,
        education=education,
        projects=projects,
        experience=experience,
        internships=internships,
        certifications=certifications,
        achievements=achievements,
        graduation_date=_extract_graduation_date(education),
        estimated_years_experience=years,
        field_confidence=confidence,
    )
    return ParseResult(fields=fields, warnings=warnings)


def parse_resume_bytes(data: bytes, extension: str) -> tuple[str, ParseResult]:
    """Extract text then parse it. Returns the text and the parse result."""
    text = extract_text(data, extension)
    result = parse_resume_text(text)
    if extension.lower() == ".pdf":
        merged = list(result.fields.links)
        for uri in extract_pdf_links(data):
            if uri not in merged:
                merged.append(uri)
        result.fields.links = merged[:10]
    return text, result


def soft_skills_from_text(text: str) -> list[str]:
    return _extract_soft_skills(text)
